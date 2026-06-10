#!/usr/bin/env python3
"""Build the 0609 ROD dataset preparation slide deck."""

from __future__ import annotations

import shutil
from pathlib import Path

import cv2
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
OUT_DIR = ROOT / "plans" / "0609_pre"
ASSET_DIR = OUT_DIR / "assets"

ROD_EXAMPLES = Path("/mnt/drive/3333_raw/ROD/example_comparisons")
VIS_ROOT = ROD_EXAMPLES / "0603_rawvis_student_teacher_dav2s_teacher_dav2l_6col_spectral_r"
RGB_ROOT = ROD_EXAMPLES / "0603_rawvis_teacher_student_night_degreen_v1_20samples"
DAV2L_ROOT = ROD_EXAMPLES / "0603_rawvis_teacher_student_night_degreen_v1_20samples_dav2l_spectral_r"
HIST_ROOT = ROD_EXAMPLES / "0603_6col_raw_rggb_student_teacher_rgb_channel_hists"

SAMPLE = "00Train_night-02510"
SAMPLE_ID = "night-02510"

LOD_ROOT = Path("/home/caq/6666_raw/0000_dataset/LOD")
LOD_PAIR_ID = "lod-0039-0040"
LOD_NORMAL_ID = "39"
LOD_DARK_ID = "40"
LOD_PSEUDO_NAME = "0039_0040.npy"
LOD_NOISE_COMPARE_NAME = "lod_0039_0040_brightness_noise_compare_cn.png"

FONT = "Noto Sans CJK SC"
TITLE = RGBColor(28, 35, 45)
BODY = RGBColor(55, 65, 81)
MUTED = RGBColor(107, 114, 128)
BLUE = RGBColor(37, 99, 235)
TEAL = RGBColor(15, 118, 110)
AMBER = RGBColor(180, 83, 9)
PURPLE = RGBColor(109, 40, 217)
GREEN = RGBColor(22, 163, 74)
LIGHT_BLUE = RGBColor(239, 246, 255)
LIGHT_TEAL = RGBColor(240, 253, 250)
LIGHT_AMBER = RGBColor(255, 251, 235)
LIGHT_PURPLE = RGBColor(245, 243, 255)
LIGHT_GREEN = RGBColor(240, 253, 244)
LINE = RGBColor(209, 213, 219)
BG = RGBColor(248, 250, 252)
WHITE = RGBColor(255, 255, 255)


def px(value: float) -> int:
    return int(round(value))


def ensure_dirs() -> None:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)


def copy_asset(src: Path, dst_name: str) -> Path:
    dst = ASSET_DIR / dst_name
    if not src.is_file():
        raise FileNotFoundError(src)
    shutil.copy2(src, dst)
    return dst


def crop_cover(src: Path, dst: Path, aspect: float) -> Path:
    img = Image.open(src).convert("RGB")
    w, h = img.size
    current = w / h
    if current > aspect:
        new_w = px(h * aspect)
        left = (w - new_w) // 2
        box = (left, 0, left + new_w, h)
    else:
        new_h = px(w / aspect)
        top = max(0, (h - new_h) // 2)
        box = (0, top, w, top + new_h)
    img.crop(box).save(dst, quality=95)
    return dst


def crop_hist_panel(src: Path, dst: Path) -> Path:
    """Crop first sample's RAW / student / teacher histogram panels from the existing contact sheet."""

    img = Image.open(src).convert("RGB")
    w, h = img.size
    # The source contact sheet is a 6-column panel; the first row contains image
    # thumbnails followed by RGB/RGGB histograms. This crop keeps only the three
    # histograms needed by this slide.
    y0, y1 = int(h * 0.111), int(h * 0.191)
    raw = img.crop((0, y0, w // 6, y1))
    student = img.crop((w // 6, y0, 2 * w // 6, y1))
    teacher = img.crop((3 * w // 6, y0, 4 * w // 6, y1))

    pad = 18
    out_w = raw.width + student.width + teacher.width + 4 * pad
    out_h = raw.height + 2 * pad
    out = Image.new("RGB", (out_w, out_h), (255, 255, 255))
    x = pad
    for panel in (raw, student, teacher):
        out.paste(ImageOps.expand(panel, border=1, fill=(218, 224, 232)), (x, pad))
        x += panel.width + pad
    out.save(dst, quality=95)
    return dst


def stretch_p1p99_rgb(rgb: np.ndarray) -> np.ndarray:
    rgb_f = rgb.astype(np.float32, copy=False)
    out_channels = []
    for channel in range(3):
        ch = rgb_f[..., channel]
        lo = float(np.percentile(ch, 1.0))
        hi = float(np.percentile(ch, 99.0))
        if hi <= lo:
            hi = lo + 1.0
        out_channels.append(np.clip((ch - lo) / (hi - lo), 0.0, 1.0))
    return np.stack(out_channels, axis=-1)


def load_lod_rgb8(src: Path) -> np.ndarray:
    bgr = cv2.imread(str(src), cv2.IMREAD_UNCHANGED)
    if bgr is None:
        raise FileNotFoundError(src)
    if bgr.ndim != 3 or bgr.shape[-1] != 3 or bgr.dtype != np.uint8:
        raise ValueError(f"Expected LOD RGB JPG HWC uint8, got {bgr.shape} {bgr.dtype}: {src}")
    return bgr[..., ::-1]


def load_lod_raw_rgb16(src: Path) -> np.ndarray:
    raw_bgr = cv2.imread(str(src), cv2.IMREAD_UNCHANGED)
    if raw_bgr is None:
        raise FileNotFoundError(src)
    if raw_bgr.ndim != 3 or raw_bgr.shape[-1] != 3 or raw_bgr.dtype != np.uint16:
        raise ValueError(f"Expected LOD RAW RGB16 PNG HWC uint16, got {raw_bgr.shape} {raw_bgr.dtype}: {src}")
    return raw_bgr[..., ::-1]


def render_lod_rgb8_p1p99(src: Path, dst: Path) -> Path:
    rgb = load_lod_rgb8(src)
    preview = stretch_p1p99_rgb(rgb)
    Image.fromarray((preview * 255.0).astype(np.uint8)).save(dst, quality=95)
    return dst


def render_lod_raw_rgb16(src: Path, dst: Path) -> Path:
    raw_rgb = load_lod_raw_rgb16(src)
    preview = stretch_p1p99_rgb(raw_rgb)
    Image.fromarray((preview * 255.0).astype(np.uint8)).save(dst, quality=95)
    return dst


def render_depth_npy(src: Path, dst: Path) -> Path:
    pred = np.load(src).astype(np.float32, copy=False)
    valid = pred[np.isfinite(pred)]
    if valid.size == 0:
        norm = np.zeros_like(pred, dtype=np.float32)
    else:
        lo = float(np.percentile(valid, 1.0))
        hi = float(np.percentile(valid, 99.0))
        if hi <= lo:
            hi = lo + 1e-6
        norm = np.clip((pred - lo) / (hi - lo), 0.0, 1.0)
    cmap = matplotlib.colormaps.get_cmap("Spectral_r")
    colored_rgb = (cmap(norm)[..., :3] * 255.0).round().astype(np.uint8)
    Image.fromarray(colored_rgb).save(dst, quality=95)
    return dst


def sample_values(values: np.ndarray, max_points: int = 220_000) -> np.ndarray:
    finite = values[np.isfinite(values)]
    if finite.size <= max_points:
        return finite
    step = max(1, finite.size // max_points)
    return finite[::step]


def compact_number(value: float) -> str:
    if abs(value) >= 1000:
        return f"{value / 1000:.1f}k"
    if abs(value) >= 10:
        return f"{value:.0f}"
    return f"{value:.2f}"


def save_compact_hist(
    dst: Path,
    series: list[tuple[str, np.ndarray, str]],
) -> Path:
    fig, ax = plt.subplots(figsize=(5.8, 0.82), dpi=220)
    ax.set_facecolor("white")
    ymax = 1.0
    for label, values, color in series:
        vals = values.astype(np.float32, copy=False)
        vals = np.clip(vals[np.isfinite(vals)], 0.0, 1.0)
        counts, edges = np.histogram(vals, bins=96, range=(0.0, 1.0))
        ymax = max(ymax, float(counts.max()))
        centers = (edges[:-1] + edges[1:]) * 0.5
        ax.plot(centers, counts, color=color, linewidth=0.78, label=label)

    ax.set_xlim(0.0, 1.0)
    ax.set_ylim(0.0, ymax * 1.05)
    ax.set_xticks([0.0, 0.5, 1.0])
    ax.set_yticks([0.0, ymax])
    ax.set_yticklabels(["0", compact_number(ymax)])
    ax.tick_params(axis="x", labelsize=4.2, length=1.4, pad=1.0, colors="#64748b")
    ax.tick_params(axis="y", labelsize=4.0, length=1.2, pad=1.0, colors="#64748b")
    ax.text(0.985, 0.84, "raw values, x in [0,1]", transform=ax.transAxes, ha="right", va="top", fontsize=4.4, color="#64748b")
    ax.legend(
        loc="upper left",
        fontsize=4.3,
        frameon=False,
        ncol=len(series),
        handlelength=1.0,
        handletextpad=0.25,
    )
    for spine in ax.spines.values():
        spine.set_color("#e2e8f0")
        spine.set_linewidth(0.45)
    fig.tight_layout(pad=0.12)
    fig.savefig(dst, facecolor="white")
    plt.close(fig)
    return dst


def render_lod_channel_hist(src: Path, dst: Path, kind: str) -> Path:
    if kind == "rgb8":
        data = load_lod_rgb8(src).astype(np.float32) / 255.0
    elif kind == "raw16":
        data = load_lod_raw_rgb16(src).astype(np.float32) / 65535.0
    else:
        raise ValueError(f"Unknown LOD histogram kind: {kind}")

    series = [
        ("R", data[..., 0].ravel(), "#dc2626"),
        ("G", data[..., 1].ravel(), "#16a34a"),
        ("B", data[..., 2].ravel(), "#2563eb"),
    ]
    return save_compact_hist(dst, series)


def build_lod_assets() -> dict[str, Path]:
    rgb_normal_src = LOD_ROOT / "RGB_normal" / f"{LOD_NORMAL_ID}.JPG"
    rgb_dark_src = LOD_ROOT / "RGB_Dark" / f"{LOD_DARK_ID}.JPG"
    raw_normal_src = LOD_ROOT / "RAW_normal" / f"{LOD_NORMAL_ID}.png"
    raw_dark_src = LOD_ROOT / "RAW_Dark" / f"{LOD_DARK_ID}.png"
    pseudo_src = LOD_ROOT / "pseudo_depth_dav2l_rgb_normal_rel_1200x800" / "00Train" / LOD_PSEUDO_NAME
    noise_compare = ASSET_DIR / LOD_NOISE_COMPARE_NAME
    if not noise_compare.is_file():
        raise FileNotFoundError(noise_compare)

    rgb_normal = render_lod_rgb8_p1p99(
        rgb_normal_src,
        ASSET_DIR / "lod_0039_rgb_normal_p1p99.jpg",
    )
    rgb_dark = render_lod_rgb8_p1p99(
        rgb_dark_src,
        ASSET_DIR / "lod_0040_rgb_dark_p1p99.jpg",
    )
    raw_normal = render_lod_raw_rgb16(
        raw_normal_src,
        ASSET_DIR / "lod_0039_raw_normal_p1p99.jpg",
    )
    raw_dark = render_lod_raw_rgb16(
        raw_dark_src,
        ASSET_DIR / "lod_0040_raw_dark_p1p99.jpg",
    )
    pseudo = render_depth_npy(
        pseudo_src,
        ASSET_DIR / "lod_0039_0040_dav2l_pseudo_depth.jpg",
    )

    tile_paths = {}
    for name, path in {
        "lod_rgb_normal_tile": rgb_normal,
        "lod_raw_normal_tile": raw_normal,
        "lod_rgb_dark_tile": rgb_dark,
        "lod_raw_dark_tile": raw_dark,
    }.items():
        tile_paths[name] = crop_cover(path, ASSET_DIR / f"{name}.jpg", aspect=1.73)
    tile_paths["lod_pseudo_tile"] = crop_cover(pseudo, ASSET_DIR / "lod_pseudo_tile.jpg", aspect=2.24)

    lod_hists = {
        "lod_rgb_normal_hist": render_lod_channel_hist(
            rgb_normal_src,
            ASSET_DIR / "lod_0039_rgb_normal_hist.png",
            "rgb8",
        ),
        "lod_raw_normal_hist": render_lod_channel_hist(
            raw_normal_src,
            ASSET_DIR / "lod_0039_raw_normal_hist.png",
            "raw16",
        ),
        "lod_rgb_dark_hist": render_lod_channel_hist(
            rgb_dark_src,
            ASSET_DIR / "lod_0040_rgb_dark_hist.png",
            "rgb8",
        ),
        "lod_raw_dark_hist": render_lod_channel_hist(
            raw_dark_src,
            ASSET_DIR / "lod_0040_raw_dark_hist.png",
            "raw16",
        ),
    }

    return {
        "lod_rgb_normal": rgb_normal,
        "lod_rgb_dark": rgb_dark,
        "lod_raw_normal": raw_normal,
        "lod_raw_dark": raw_dark,
        "lod_pseudo": pseudo,
        "lod_noise_compare": noise_compare,
        **tile_paths,
        **lod_hists,
    }


def build_assets() -> dict[str, Path]:
    ensure_dirs()
    lod_assets = build_lod_assets()
    raw_vis = copy_asset(
        VIS_ROOT / "raw_visualized_default" / "01_00Train_night-02510_raw_visualized_default.png",
        "night-02510_raw_visualized.png",
    )
    student = copy_asset(
        RGB_ROOT / "student_dark_degreen_v1" / f"{SAMPLE}.png",
        "night-02510_student_dark_degreen_v1.png",
    )
    teacher = copy_asset(
        RGB_ROOT / "teacher_bright_degreen_v1" / f"{SAMPLE}.png",
        "night-02510_teacher_bright_degreen_v1.png",
    )
    pseudo = copy_asset(
        DAV2L_ROOT
        / "predictions_spectral_r"
        / "01_00Train_night-02510_teacher_bright_degreen_v1_dav2l_spectral_r.png",
        "night-02510_teacher_dav2l_pseudo_depth.png",
    )
    hist_panel = crop_hist_panel(
        HIST_ROOT
        / "panels"
        / "0603_night_6col_raw_rggb_student_teacher_rgb_channel_hists_page1.jpg",
        ASSET_DIR / "night-02510_existing_channel_hist_crops.jpg",
    )

    tiles = {}
    for name, path in {
        "raw_tile": raw_vis,
        "student_tile": student,
        "teacher_tile": teacher,
        "pseudo_tile": pseudo,
    }.items():
        tiles[name] = crop_cover(path, ASSET_DIR / f"{name}.jpg", aspect=1.45)

    return {
        "raw": raw_vis,
        "student": student,
        "teacher": teacher,
        "pseudo": pseudo,
        "hist_panel": hist_panel,
        **tiles,
        **lod_assets,
    }


def set_fill(shape, color: RGBColor, transparency: int = 0) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.fill.transparency = transparency
    shape.line.color.rgb = LINE
    shape.line.width = Pt(0.7)


def add_textbox(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    size: float = 12,
    color: RGBColor = BODY,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.margin_left = Inches(0.03)
    box.text_frame.margin_right = Inches(0.03)
    box.text_frame.margin_top = Inches(0.02)
    box.text_frame.margin_bottom = Inches(0.02)
    box.text_frame.vertical_anchor = valign
    box.text_frame.clear()
    p = box.text_frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_rich_text(slide, x: float, y: float, w: float, h: float, lines: list[tuple[str, float, RGBColor, bool]]) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.02)
    tf.word_wrap = True
    tf.clear()
    for i, (text, size, color, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(2)
        r = p.add_run()
        r.text = text
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold


def add_card(slide, x: float, y: float, w: float, h: float, fill: RGBColor = WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.adjustments[0] = 0.08
    set_fill(shape, fill)
    return shape


def add_tile(slide, path: Path, x: float, y: float, w: float, h: float, title: str, accent: RGBColor) -> None:
    add_card(slide, x, y, w, h, WHITE)
    pic = slide.shapes.add_picture(str(path), Inches(x + 0.04), Inches(y + 0.30), width=Inches(w - 0.08), height=Inches(h - 0.46))
    pic.line.color.rgb = RGBColor(226, 232, 240)
    pic.line.width = Pt(0.5)
    add_textbox(slide, x + 0.10, y + 0.08, w - 0.20, 0.18, title, size=8.5, color=accent, bold=True)


def add_lod_tile_with_hist(
    slide,
    image_path: Path,
    hist_path: Path,
    x: float,
    y: float,
    w: float,
    h: float,
    title: str,
    accent: RGBColor,
) -> None:
    add_card(slide, x, y, w, h, WHITE)
    add_textbox(slide, x + 0.10, y + 0.08, w - 0.20, 0.18, title, size=8.0, color=accent, bold=True)

    hist_h = 0.36
    hist_y = y + h - hist_h - 0.07
    pic_y = y + 0.30
    pic_h = hist_y - pic_y - 0.06
    pic = slide.shapes.add_picture(
        str(image_path),
        Inches(x + 0.04),
        Inches(pic_y),
        width=Inches(w - 0.08),
        height=Inches(pic_h),
    )
    pic.line.color.rgb = RGBColor(226, 232, 240)
    pic.line.width = Pt(0.5)

    hist = slide.shapes.add_picture(
        str(hist_path),
        Inches(x + 0.08),
        Inches(hist_y),
        width=Inches(w - 0.16),
        height=Inches(hist_h),
    )
    hist.line.color.rgb = RGBColor(226, 232, 240)
    hist.line.width = Pt(0.35)


def add_arrow(slide, x1: float, y1: float, x2: float, y2: float, color: RGBColor = RGBColor(148, 163, 184)) -> None:
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = color
    line.line.width = Pt(1.2)
    line.line.end_arrowhead = True


def add_line(slide, x1: float, y1: float, x2: float, y2: float, color: RGBColor, width: float = 0.6) -> None:
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = color
    line.line.width = Pt(width)


def add_rect(slide, x: float, y: float, w: float, h: float, fill: RGBColor, transparency: int = 0):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.fill.transparency = transparency
    shape.line.fill.background()
    try:
        shape.shadow.inherit = False
    except AttributeError:
        pass
    return shape


def add_flow_box(slide, x: float, y: float, w: float, h: float, title: str, body: str, fill: RGBColor, accent: RGBColor) -> None:
    add_card(slide, x, y, w, h, fill)
    add_textbox(slide, x + 0.10, y + 0.08, w - 0.20, 0.20, title, size=9.4, color=accent, bold=True)
    add_textbox(slide, x + 0.10, y + 0.35, w - 0.20, h - 0.42, body, size=7.6, color=BODY)


def add_lod_data_display_slide(prs: Presentation, assets: dict[str, Path]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG

    add_textbox(
        slide,
        0.42,
        0.20,
        8.70,
        0.45,
        "LOD 数据展示：normal / dark × RGB / RAW",
        size=20.5,
        color=TITLE,
        bold=True,
    )
    add_textbox(
        slide,
        9.55,
        0.28,
        3.25,
        0.30,
        "阶段性汇报 · 2026-06-09",
        size=9.5,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )

    add_card(slide, 0.42, 0.80, 3.00, 1.72, WHITE)
    add_textbox(slide, 0.58, 0.95, 2.68, 0.24, "数据本体：四个目录", size=12.2, color=BLUE, bold=True)
    data_lines = [
        "`RGB_normal/`：正常曝光 JPG",
        "`RGB_Dark/`：低光 JPG",
        "`RAW_normal/`：正常曝光 RAW PNG",
        "`RAW_Dark/`：低光 RAW PNG",
        "每个 pair：normal_id 与 dark_id 成对",
    ]
    add_rich_text(slide, 0.58, 1.25, 2.68, 1.02, [(f"• {s}", 7.55, BODY, False) for s in data_lines])

    add_card(slide, 0.42, 2.76, 3.00, 1.28, WHITE)
    add_textbox(slide, 0.58, 2.91, 2.68, 0.24, "RAW 物理格式", size=12.0, color=TEAL, bold=True)
    raw_lines = [
        "3 通道 16-bit PNG，非 Bayer",
        "OpenCV 读入 BGR，模型侧重排为 RGB",
        "训练归一化：uint16 / 65535",
        "本页 RGB/RAW 均做 p1/p99 拉伸可视化",
    ]
    add_rich_text(slide, 0.58, 3.19, 2.68, 0.66, [(f"• {s}", 7.1, BODY, False) for s in raw_lines])

    add_card(slide, 0.42, 4.30, 3.00, 2.18, WHITE)
    add_textbox(slide, 0.58, 4.46, 2.68, 0.24, "当前训练用法", size=12.0, color=AMBER, bold=True)
    train_lines = [
        "teacher：RGB_normal → DAv2-L → 2D pseudo depth",
        "RGB 线：RGB_Dark 作为 DAv2-S 输入",
        "RAW 线：RAW_Dark 作为 RAW 前端输入",
        "RAW_normal：只作为 normal/raw 对比参照",
        "公平主对比：同一 dark 场景下 RGB_Dark vs RAW_Dark",
    ]
    add_rich_text(slide, 0.58, 4.76, 2.68, 1.34, [(f"• {s}", 7.15, BODY, False) for s in train_lines])

    # Four-source visual matrix.
    tile_w = 2.22
    tile_h = 2.03
    x_left = 3.72
    x_right = 6.16
    y_top = 0.98
    y_bottom = 3.16
    add_lod_tile_with_hist(
        slide,
        assets["lod_rgb_normal_tile"],
        assets["lod_rgb_normal_hist"],
        x_left,
        y_top,
        tile_w,
        tile_h,
        "RGB_normal · P1/P99",
        BLUE,
    )
    add_lod_tile_with_hist(
        slide,
        assets["lod_raw_normal_tile"],
        assets["lod_raw_normal_hist"],
        x_right,
        y_top,
        tile_w,
        tile_h,
        "RAW_normal · P1/P99",
        TEAL,
    )
    add_lod_tile_with_hist(
        slide,
        assets["lod_rgb_dark_tile"],
        assets["lod_rgb_dark_hist"],
        x_left,
        y_bottom,
        tile_w,
        tile_h,
        "RGB_Dark · RGB 线输入",
        AMBER,
    )
    add_lod_tile_with_hist(
        slide,
        assets["lod_raw_dark_tile"],
        assets["lod_raw_dark_hist"],
        x_right,
        y_bottom,
        tile_w,
        tile_h,
        "RAW_Dark · RAW 线输入",
        PURPLE,
    )
    add_textbox(slide, 3.72, 0.74, 4.66, 0.18, f"示例 pair：{LOD_PAIR_ID}（normal={LOD_NORMAL_ID}, dark={LOD_DARK_ID}）", size=7.4, color=MUTED)

    add_tile(
        slide,
        assets["lod_pseudo_tile"],
        9.06,
        0.98,
        3.20,
        2.18,
        "DAv2-L pseudo depth · Spectral_r",
        PURPLE,
    )

    add_flow_box(
        slide,
        9.06,
        3.40,
        1.48,
        0.76,
        "teacher",
        "RGB_normal\nDAv2-L",
        LIGHT_PURPLE,
        PURPLE,
    )
    add_flow_box(
        slide,
        10.82,
        3.40,
        1.44,
        0.76,
        "target",
        "inverse-relative\npseudo depth",
        WHITE,
        PURPLE,
    )
    add_arrow(slide, 10.54, 3.78, 10.82, 3.78, PURPLE)

    add_flow_box(
        slide,
        9.06,
        4.44,
        1.48,
        0.82,
        "RGB student",
        "RGB_Dark\nDAv2-S",
        LIGHT_AMBER,
        AMBER,
    )
    add_flow_box(
        slide,
        10.82,
        4.44,
        1.44,
        0.82,
        "RAW student",
        "RAW_Dark\nRAW front-end",
        LIGHT_TEAL,
        TEAL,
    )
    add_textbox(
        slide,
        9.08,
        5.31,
        3.12,
        0.20,
        "两条 student 线共享同一 Spectral_r pseudo label；差别只在 dark 输入域。",
        size=6.35,
        color=BODY,
    )

    add_card(slide, 3.72, 5.58, 8.54, 0.88, LIGHT_BLUE)
    add_textbox(slide, 3.95, 5.71, 8.05, 0.20, "对比口径", size=10.0, color=BLUE, bold=True)
    compare_text = (
        "主对比：RGB_Dark vs RAW_Dark\n"
        "同一低光场景、同一 crop、同一 RGB_normal→DAv2-L pseudo GT。\n"
        "RAW_normal 仅作为 normal/raw 可视化与诊断参照，不进入当前主线输入。"
    )
    add_textbox(slide, 3.95, 5.96, 8.04, 0.38, compare_text, size=6.65, color=BODY)

    add_textbox(
        slide,
        0.46,
        7.13,
        12.42,
        0.20,
        "代码依据：finetune_stf/dataset/lod_true.py；数据依据：/home/caq/6666_raw/0000_dataset/LOD/{RGB_normal,RGB_Dark,RAW_normal,RAW_Dark}/",
        size=6.7,
        color=MUTED,
    )


def add_lod_noise_analysis_slide(prs: Presentation, assets: dict[str, Path]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG

    img = Image.open(assets["lod_noise_compare"])
    aspect = img.width / img.height
    h = 7.04
    w = h * aspect
    x = (13.333 - w) / 2.0
    y = 0.20
    pic = slide.shapes.add_picture(
        str(assets["lod_noise_compare"]),
        Inches(x),
        Inches(y),
        width=Inches(w),
        height=Inches(h),
    )
    pic.line.color.rgb = RGBColor(226, 232, 240)
    pic.line.width = Pt(0.5)


def add_rod_result_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG

    add_textbox(
        slide,
        0.42,
        0.20,
        8.70,
        0.45,
        "ROD 结果对比：以 d1 为主指标",
        size=20.5,
        color=TITLE,
        bold=True,
    )
    add_textbox(
        slide,
        9.55,
        0.28,
        3.25,
        0.30,
        "阶段性汇报 · 2026-06-09",
        size=9.5,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )

    rows = [
        ("init RGB D0", "official DAv2-S init", 0.7617, "-", 13.5764, "RGB", MUTED),
        ("0604_0752", "RGB decoder-only", 0.8015, "e7", 9.1288, "RGB", BLUE),
        ("0606_1348 R1", "RAW-RAM decoder repl", 0.8412, "e6", 5.5653, "RAW", TEAL),
        ("0605_0139", "RGB LoRA tap + decoder", 0.8251, "e7", 8.3252, "RGB", BLUE),
        ("0606_1348 R2", "RAW-RAM LoRA repl", 0.8437, "e3", 6.8383, "RAW", TEAL),
        ("0605_0730", "RGB full backbone LLRD", 0.8483, "e8", 5.8880, "RGB", GREEN),
        ("0606_0330", "RAW-RAM full backbone", 0.8700, "e9", 4.2575, "RAW", PURPLE),
        ("0608_1444", "RAW-RAM LoRA + bridge/adapter", 0.8517, "e4", 5.6001, "RAW", AMBER),
    ]

    add_card(slide, 0.42, 0.82, 8.62, 5.86, WHITE)
    add_textbox(slide, 0.62, 0.98, 4.35, 0.24, "best checkpoint 指标表（d1 为主）", size=12.0, color=TITLE, bold=True)
    add_textbox(slide, 5.30, 0.99, 1.08, 0.20, "d1 best", size=8.9, color=BLUE, bold=True, align=PP_ALIGN.RIGHT)
    add_textbox(slide, 6.40, 0.99, 0.48, 0.20, "epoch", size=7.7, color=MUTED, align=PP_ALIGN.CENTER)
    add_textbox(slide, 7.12, 0.99, 1.42, 0.20, "abs_rel best", size=8.9, color=MUTED, align=PP_ALIGN.RIGHT)

    row_y0 = 1.34
    row_gap = 0.58
    add_line(slide, 0.62, 1.23, 8.72, 1.23, RGBColor(226, 232, 240), width=0.7)

    for idx, (run, method, d1, epoch, abs_rel, domain, accent) in enumerate(rows):
        y = row_y0 + idx * row_gap
        is_best = run == "0606_0330"
        add_line(slide, 0.62, y + 0.42, 8.72, y + 0.42, RGBColor(241, 245, 249), width=0.45)
        add_textbox(slide, 0.72, y - 0.02, 1.30, 0.20, run, size=8.6, color=accent)
        add_textbox(slide, 2.06, y - 0.02, 0.44, 0.20, domain, size=6.9, color=MUTED)
        add_textbox(slide, 2.54, y - 0.02, 2.22, 0.20, method, size=7.15, color=BODY)

        add_textbox(slide, 5.16, y - 0.05, 1.24, 0.28, f"{d1:.4f}", size=12.8, color=TITLE, align=PP_ALIGN.RIGHT)
        add_textbox(slide, 6.46, y - 0.01, 0.42, 0.20, epoch, size=7.7, color=MUTED, align=PP_ALIGN.CENTER)
        add_textbox(slide, 7.10, y - 0.04, 1.44, 0.26, f"{abs_rel:.4f}", size=10.7, color=TITLE if is_best else BODY, align=PP_ALIGN.RIGHT)

    add_card(slide, 9.28, 0.82, 3.62, 1.52, LIGHT_PURPLE)
    add_textbox(slide, 9.48, 0.99, 3.22, 0.24, "当前最强", size=12.4, color=PURPLE, bold=True)
    add_textbox(slide, 9.48, 1.29, 3.18, 0.70, "0606_0330：RAW-RAM + full backbone\nd1=0.8700，abs_rel=4.2575", size=8.0, color=BODY)

    add_card(slide, 9.28, 2.58, 3.62, 1.28, LIGHT_GREEN)
    add_textbox(slide, 9.48, 2.74, 3.22, 0.22, "RGB 路径上限", size=11.3, color=GREEN, bold=True)
    add_textbox(slide, 9.48, 3.03, 3.18, 0.52, "0605_0730：full backbone\nd1=0.8483，高于 RGB decoder / LoRA", size=7.45, color=BODY)

    add_card(slide, 9.28, 4.04, 3.62, 1.28, LIGHT_AMBER)
    add_textbox(slide, 9.48, 4.20, 3.22, 0.22, "Bridge / adapter", size=11.3, color=AMBER, bold=True)
    add_textbox(slide, 9.48, 4.49, 3.18, 0.52, "0608_1444：LoRA + bridge/adapter\nd1=0.8517，高于两个 0606_1348", size=7.35, color=BODY)

    add_card(slide, 9.28, 5.54, 3.62, 1.14, LIGHT_BLUE)
    add_textbox(slide, 9.48, 5.68, 3.22, 0.22, "读数口径", size=11.0, color=BLUE, bold=True)
    add_textbox(slide, 9.48, 5.94, 3.18, 0.48, "init 为 RGB D0；RAW-RAM init d1≈0.270，\n不作为公共 baseline。", size=7.0, color=BODY)

    add_textbox(
        slide,
        0.46,
        7.05,
        12.42,
        0.30,
        "备注：ROD 指标来自 DAv2-L pseudo inverse-relative label；abs_rel 不是 metric-depth 绝对误差，数值只适合同一 ROD-night eval protocol 内作辅助趋势参考。",
        size=6.85,
        color=MUTED,
    )


def add_cross_eval_matrix(
    slide,
    x: float,
    y: float,
    w: float,
    title: str,
    rows: list[tuple[str, float, float]],
    recovery_text: str,
    accent: RGBColor,
) -> None:
    add_textbox(slide, x, y, w, 0.22, title, size=8.2, color=accent, bold=True)
    add_textbox(slide, x + 1.28, y + 0.30, 0.78, 0.18, "I_dark", size=6.2, color=MUTED, align=PP_ALIGN.CENTER)
    add_textbox(slide, x + 2.18, y + 0.30, 0.86, 0.18, "I_normal", size=6.2, color=MUTED, align=PP_ALIGN.CENTER)
    add_line(slide, x, y + 0.53, x + w, y + 0.53, RGBColor(226, 232, 240), width=0.55)

    row_y = y + 0.60
    row_h = 0.34
    for idx, (label, dark_d1, normal_d1) in enumerate(rows):
        yy = row_y + idx * row_h
        add_textbox(slide, x, yy + 0.03, 1.16, 0.18, label, size=6.0, color=BODY, bold=True)
        add_textbox(slide, x + 1.28, yy, 0.78, 0.24, f"{dark_d1:.4f}", size=8.6, color=TITLE, align=PP_ALIGN.CENTER)
        add_textbox(slide, x + 2.18, yy, 0.86, 0.24, f"{normal_d1:.4f}", size=8.6, color=TITLE, align=PP_ALIGN.CENTER)
        add_line(slide, x, yy + 0.29, x + w, yy + 0.29, RGBColor(241, 245, 249), width=0.45)

    add_textbox(slide, x, y + 1.36, w, 0.26, recovery_text, size=5.8, color=BODY)


def add_lod_cross_eval_panel(slide) -> None:
    add_card(slide, 9.28, 0.82, 3.62, 5.86, WHITE)
    add_textbox(slide, 9.48, 0.99, 3.22, 0.24, "RAW Dark/Normal cross-eval", size=11.6, color=PURPLE, bold=True)
    add_textbox(slide, 9.48, 1.25, 3.22, 0.18, "strict variant · D1 on 01Valid / 112 samples", size=6.15, color=MUTED)

    add_cross_eval_matrix(
        slide,
        9.48,
        1.62,
        3.12,
        "Canonical LoRA pair",
        [
            ("C_dark", 0.845279, 0.880533),
            ("C_normal", 0.816972, 0.898946),
        ],
        "G=0.053667 · R_DN=65.7% · ΔDN=+0.035253 · ΔND=-0.028307",
        TEAL,
    )

    add_cross_eval_matrix(
        slide,
        9.48,
        3.42,
        3.12,
        "W0 decoder-only sanity",
        [
            ("C_dark_W0", 0.829694, 0.875950),
            ("C_normal_W0", 0.792805, 0.887073),
        ],
        "G=0.057380 · R_DN=80.6% · ΔDN=+0.046256 · ΔND=-0.036889",
        AMBER,
    )

    add_line(slide, 9.48, 5.24, 12.60, 5.24, RGBColor(226, 232, 240), width=0.55)
    add_textbox(slide, 9.48, 5.42, 3.12, 0.22, "读法", size=8.6, color=BLUE, bold=True)
    takeaways = (
        "C_dark(I_normal) recovers most, but not all, of the canonical gap.\n"
        "C_normal(I_dark) < C_dark(I_dark), so normal checkpoint is not robust to RAW_dark."
    )
    add_textbox(slide, 9.48, 5.68, 3.12, 0.48, takeaways, size=5.9, color=BODY)
    add_textbox(slide, 9.48, 6.36, 3.12, 0.16, "source: plans/0609/lod_raw_cross_eval_combined_report.md", size=5.1, color=MUTED)


def add_lod_result_slide(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG

    add_textbox(
        slide,
        0.42,
        0.20,
        8.85,
        0.45,
        "LOD 结果对比：以 d1 为主，silog 为 error 辅助",
        size=20.0,
        color=TITLE,
        bold=True,
    )
    add_textbox(
        slide,
        9.55,
        0.28,
        3.25,
        0.30,
        "阶段性汇报 · 2026-06-09",
        size=9.5,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )

    rows = [
        ("init RGB D0", "RGB_Dark init, DAv2-S", 0.8245, "-", 0.4593, "RGB", MUTED),
        ("0608_1509", "RGB_Dark decoder baseline", 0.8351, "e30", 0.6062, "RGB", BLUE),
        ("0608_1739", "RAW_Dark RGB16 decoder", 0.8299, "e36", 0.5397, "RAW_D", TEAL),
        ("0608_2246", "RGB_Dark LoRA tap", 0.8555, "e22", 0.5851, "RGB", GREEN),
        ("0608_2026", "RAW_Dark RGB16 LoRA tap", 0.8451, "e37", 0.5523, "RAW_D", TEAL),
        ("0608_2059", "RAW_Dark RGB16 LoRA + bridge/FA", 0.8510, "e27", 0.5690, "RAW_D", PURPLE),
        ("0609_0044", "RAW_Normal RGB16 LoRA tap", 0.8990, "e36", 0.4912, "RAW_N", AMBER),
        ("0609_0117", "RAW_Normal RGB16 LoRA + bridge/FA", 0.8997, "e34", 0.5417, "RAW_N", PURPLE),
    ]

    add_card(slide, 0.42, 0.82, 8.62, 5.86, WHITE)
    add_textbox(slide, 0.62, 0.98, 4.35, 0.24, "best checkpoint 指标表（d1 为主）", size=12.0, color=TITLE, bold=True)
    add_textbox(slide, 5.30, 0.99, 1.08, 0.20, "d1 best", size=8.9, color=BLUE, bold=True, align=PP_ALIGN.RIGHT)
    add_textbox(slide, 6.40, 0.99, 0.48, 0.20, "epoch", size=7.7, color=MUTED, align=PP_ALIGN.CENTER)
    add_textbox(slide, 7.05, 0.99, 1.50, 0.20, "silog @best", size=8.9, color=MUTED, align=PP_ALIGN.RIGHT)

    row_y0 = 1.34
    row_gap = 0.58
    add_line(slide, 0.62, 1.23, 8.72, 1.23, RGBColor(226, 232, 240), width=0.7)

    for idx, (run, method, d1, epoch, silog, domain, accent) in enumerate(rows):
        y = row_y0 + idx * row_gap
        is_best = run == "0609_0117"
        add_line(slide, 0.62, y + 0.42, 8.72, y + 0.42, RGBColor(241, 245, 249), width=0.45)
        add_textbox(slide, 0.72, y - 0.02, 1.30, 0.20, run, size=8.6, color=accent)
        add_textbox(slide, 2.05, y - 0.02, 0.52, 0.20, domain, size=6.9, color=MUTED)
        add_textbox(slide, 2.60, y - 0.02, 2.25, 0.20, method, size=7.05, color=BODY)

        add_textbox(slide, 5.16, y - 0.05, 1.24, 0.28, f"{d1:.4f}", size=12.8, color=TITLE, align=PP_ALIGN.RIGHT)
        add_textbox(slide, 6.46, y - 0.01, 0.42, 0.20, epoch, size=7.7, color=MUTED, align=PP_ALIGN.CENTER)
        add_textbox(slide, 7.10, y - 0.04, 1.44, 0.26, f"{silog:.4f}", size=10.7, color=TITLE if is_best else BODY, align=PP_ALIGN.RIGHT)

    add_lod_cross_eval_panel(slide)

    add_textbox(
        slide,
        0.46,
        7.05,
        12.42,
        0.30,
        "备注：LOD 指标来自 DAv2-L pseudo inverse-relative label；本页不展示 abs_rel，使用 silog/d1 作为同一 True LOD eval protocol 内的 proxy 口径。RAW init d1≈0.370，不是 RGB init 的公共 baseline。",
        size=6.85,
        color=MUTED,
    )


def build_pptx(assets: dict[str, Path]) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    add_lod_data_display_slide(prs, assets)
    add_lod_noise_analysis_slide(prs, assets)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG

    add_textbox(
        slide,
        0.42,
        0.20,
        8.60,
        0.45,
        "ROD 数据集准备：RAW24 → student RGB / teacher RGB / pseudo GT",
        size=20.5,
        color=TITLE,
        bold=True,
    )
    add_textbox(
        slide,
        9.55,
        0.28,
        3.25,
        0.30,
        "阶段性汇报 · 2026-06-09",
        size=9.5,
        color=MUTED,
        align=PP_ALIGN.RIGHT,
    )

    add_card(slide, 0.42, 0.80, 3.00, 2.04, WHITE)
    add_textbox(slide, 0.58, 0.96, 2.68, 0.26, "官方 RAW 数据格式", size=12.2, color=BLUE, bold=True)
    raw_lines = [
        "文件：`.raw` packed little-endian 24-bit",
        "Bayer：true RGGB，1856×2880",
        "解包：b0 + 256·b1 + 65536·b2",
        "归一化：raw = raw24 / (2^24 - 1)",
        "训练尺寸：native 928×1440，crop 512×960",
    ]
    add_rich_text(slide, 0.58, 1.28, 2.65, 1.33, [(f"• {s}", 8.1, BODY, False) for s in raw_lines])

    add_card(slide, 0.42, 3.02, 3.00, 1.28, WHITE)
    add_textbox(slide, 0.58, 3.16, 2.68, 0.22, "核心转换公式", size=12.0, color=TEAL, bold=True)
    formula = (
        "raw = unpack24(.raw) / (2^24 - 1)\n"
        "RAW4 = [R, Gr, Gb, B]\n"
        "baseRGB = [R, (Gr+Gb)/2, B]\n"
        "RGB8 = uint8(255 · x_gamma)"
    )
    add_textbox(slide, 0.58, 3.45, 2.68, 0.68, formula, size=7.0, color=BODY)

    add_card(slide, 0.42, 4.48, 3.00, 2.12, WHITE)
    add_textbox(slide, 0.58, 4.62, 2.68, 0.22, "两条 RGB 渲染", size=11.8, color=AMBER, bold=True)
    rgb_lines = [
        "处理顺序：x0 = baseRGB",
        "g：逐通道增益，x1 = clip(x0·g, 0, 1)",
        "WP：取 x1 的 wp 百分位白点 P_wp，x2 = clip(x1/P_wp, 0, 1)",
        "gamma：x_gamma = x2^γ；γ<1 会提亮，γ越小越亮",
        "studentRGB：wp=99.9, γ=0.9，保留低光外观，作为 DAv2-S RGB 输入",
        "teacherRGB：wp=99.5, γ=0.454545，更亮，过 DAv2-L 生成 pseudo GT",
        "两者同用 g=(1.08,0.95,1.10) degreen gains",
    ]
    add_rich_text(slide, 0.58, 4.90, 2.65, 1.43, [(f"• {s}", 6.25, BODY, False) for s in rgb_lines])

    # Pipeline center
    add_flow_box(
        slide,
        3.72,
        1.05,
        1.52,
        0.76,
        "RAW24",
        "官方 .raw\n24-bit Bayer",
        LIGHT_BLUE,
        BLUE,
    )
    add_flow_box(
        slide,
        5.58,
        0.88,
        1.70,
        0.92,
        "unpack + pack",
        "归一化 full Bayer\n打包 RAW4",
        WHITE,
        TEAL,
    )
    add_flow_box(
        slide,
        7.60,
        0.78,
        1.90,
        1.02,
        "studentRGB",
        "dark ISP render\nRGB 分支输入",
        LIGHT_TEAL,
        TEAL,
    )
    add_flow_box(
        slide,
        7.60,
        2.05,
        1.90,
        1.02,
        "teacherRGB",
        "bright ISP render\nteacher 输入",
        LIGHT_AMBER,
        AMBER,
    )
    add_flow_box(
        slide,
        9.94,
        2.05,
        1.58,
        1.02,
        "DAv2-L",
        "frozen large\ninput_size=924",
        LIGHT_PURPLE,
        PURPLE,
    )
    add_flow_box(
        slide,
        11.76,
        2.05,
        1.18,
        1.02,
        "pseudo GT",
        "inverse\nrelative",
        WHITE,
        PURPLE,
    )
    add_arrow(slide, 5.24, 1.43, 5.58, 1.43)
    add_arrow(slide, 7.28, 1.34, 7.60, 1.34)
    add_arrow(slide, 7.06, 1.56, 7.60, 2.56)
    add_arrow(slide, 9.50, 2.56, 9.94, 2.56)
    add_arrow(slide, 11.52, 2.56, 11.76, 2.56)
    add_textbox(slide, 3.72, 1.90, 3.12, 0.24, "RAW 线可直接用 RAW4；RGB 对照线只吃 studentRGB", size=7.3, color=MUTED)

    # Visual tiles
    tile_y = 3.42
    tile_w = 2.15
    tile_h = 1.68
    x0 = 3.72
    gap = 0.18
    add_tile(slide, assets["raw_tile"], x0, tile_y, tile_w, tile_h, "RAW 可视化", BLUE)
    add_tile(slide, assets["student_tile"], x0 + (tile_w + gap), tile_y, tile_w, tile_h, "student RGB", TEAL)
    add_tile(slide, assets["teacher_tile"], x0 + 2 * (tile_w + gap), tile_y, tile_w, tile_h, "teacher RGB", AMBER)
    add_tile(slide, assets["pseudo_tile"], x0 + 3 * (tile_w + gap), tile_y, tile_w, tile_h, "DAv2-L pseudo GT", PURPLE)

    # Distribution area
    add_card(slide, 3.72, 5.28, 9.22, 1.62, WHITE)
    add_textbox(slide, 3.92, 5.43, 8.80, 0.20, "现成通道直方图裁剪：RAW / student / teacher", size=9.4, color=BODY, bold=True)
    slide.shapes.add_picture(str(assets["hist_panel"]), Inches(3.98), Inches(5.73), width=Inches(8.38), height=Inches(1.02))

    # Footer source note
    add_textbox(
        slide,
        0.46,
        7.13,
        12.42,
        0.20,
        "代码依据：finetune_stf/dataset/rod_raw_rgb.py · rod_raw.py · build_rod_night_teacher_labels.py；素材依据：/mnt/drive/3333_raw/ROD/example_comparisons/",
        size=6.7,
        color=MUTED,
    )

    add_rod_result_slide(prs)
    add_lod_result_slide(prs)

    pptx_path = OUT_DIR / "0609_rod_dataset_preparation_page1.pptx"
    prs.save(pptx_path)
    return pptx_path


def main() -> None:
    assets = build_assets()
    pptx_path = build_pptx(assets)
    print(pptx_path)


if __name__ == "__main__":
    main()
